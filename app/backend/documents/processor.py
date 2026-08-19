"""Clinical document parsing for DOCX/PDF — adapted from Healthcare.ipynb Week 2
(ClinicalDocumentProcessor, cells 32-35). Tracing is not baked in here; callers
wrap this with an ObservabilityPort span if they want an audit trail.
"""
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document as DocxDocument


class DocumentProcessor:
    """Parses .docx/.pdf clinical documents into paragraphs + tables + flattened text."""

    SUPPORTED_FORMATS = (".docx", ".pdf")

    def load_docx(self, file_path: Path) -> dict[str, Any] | None:
        try:
            doc = DocxDocument(str(file_path))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            tables: list[list[list[str]]] = []
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    if any(row_data):
                        table_data.append(row_data)
                if table_data:
                    tables.append(table_data)

            return {
                "filename": file_path.name,
                "format": "docx",
                "paragraphs": paragraphs,
                "tables": tables,
                "full_text": "\n".join(paragraphs),
                "table_text": self._tables_to_text(tables),
            }
        except Exception as exc:
            print(f"Error loading {file_path.name}: {exc}")
            return None

    def load_pdf(self, file_path: Path) -> dict[str, Any] | None:
        try:
            text_content: list[str] = []
            tables: list[list[list[str]]] = []

            with pdfplumber.open(str(file_path)) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_content.append(text)
                    page_tables = page.extract_tables()
                    if page_tables:
                        tables.extend(page_tables)

            return {
                "filename": file_path.name,
                "format": "pdf",
                "paragraphs": text_content,
                "tables": tables,
                "full_text": "\n".join(text_content),
                "table_text": self._tables_to_text(tables),
            }
        except Exception as exc:
            print(f"Error loading PDF {file_path.name}: {exc}")
            return None

    def _tables_to_text(self, tables: list) -> str:
        text_parts = []
        for i, table in enumerate(tables):
            text_parts.append(f"\nTable {i + 1}:")
            for row in table:
                if isinstance(row, list):
                    text_parts.append(" | ".join(str(cell) for cell in row if cell))
        return "\n".join(text_parts)

    def process_file(self, file_path: Path) -> dict[str, Any] | None:
        suffix = file_path.suffix.lower()
        if suffix == ".docx":
            return self.load_docx(file_path)
        if suffix == ".pdf":
            return self.load_pdf(file_path)
        print(f"Unsupported format: {suffix}")
        return None
